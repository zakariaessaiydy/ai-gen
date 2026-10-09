#!/usr/bin/env python3
"""gen_kids_calendar.py — build TINY-SPARKS-CALENDAR.xlsx from publishing/ideas/<niche>.txt
(100 ideas per niche, one per line). Rotation, titles, ages, formats, rig checks: see below.
Re-run after editing the idea lists:  python tools/gen_kids_calendar.py --force   (needs openpyxl)
WARNING: re-running overwrites statuses/stats typed into the xlsx — export them first."""
import datetime as dt, os, re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

D = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'publishing', 'ideas') + '/'
NICHES = [  # key, label, emoji, skill, file
    ('stories', 'Educational Stories', '🧠', 'kids-stories'), ('numbers', 'Numbers & Math', '🔢', 'kids-numbers'),
    ('words', 'English/French Words', '🔤', 'kids-words'), ('animals', 'Animal Stories', '🦁', 'kids-animals'),
    ('science', 'Science for Kids', '🚀', 'kids-science'), ('puzzles', 'Puzzles & Problem-Solving', '🧩', 'kids-puzzles'),
    ('morals', 'Moral Stories', '📚', 'kids-morals'), ('songs', 'Educational Songs', '🎵', 'kids-songs')]
IDEAS = {k: [l.strip() for l in open(D + k + '.txt', encoding='utf-8') if l.strip()] for k, *_ in NICHES}
LAB = {k: (l, e, s) for k, l, e, s in NICHES}
assert all(len(v) == 100 for v in IDEAS.values())

# ── rotation: weeks 1-100 = Mon story/moral (alternating weeks), Tue math, Wed words, Thu animals,
#    Fri science, Sat puzzles, Sun songs; then 100 days alternating the remaining stories/morals.
THEME = {0: 'Story Monday', 1: 'Math Tuesday', 2: 'Word Wednesday', 3: 'Animal Thursday', 4: 'Science Friday', 5: 'Puzzle Saturday', 6: 'Song Sunday'}
slots = []  # (day#, theme, niche, idx)
ptr = {k: 0 for k in IDEAS}
day = 1
for w in range(1, 101):
    for wd in range(7):
        if wd == 0: k = 'stories' if w % 2 == 1 else 'morals'
        else: k = ['numbers', 'words', 'animals', 'science', 'puzzles', 'songs'][wd - 1]
        slots.append((day, THEME[wd], k, ptr[k])); ptr[k] += 1; day += 1
for i in range(100):
    k = 'stories' if i % 2 == 0 else 'morals'
    slots.append((day, 'Story Season', k, ptr[k])); ptr[k] += 1; day += 1
assert all(v == 100 for v in ptr.values()) and len(slots) == 800

# ── per-idea enrichment ──
CAST_MAP = [(r'\b(Tommy|Max|Alex)\b', 'Leo'), (r'\b(Mia|Emma|Lily|Sara)\b', 'Mila'), (r'\bBenny\b', 'Bobo')]
def yt_title(k, idea):
    t = idea
    if k == 'stories':
        for pat, rep in CAST_MAP: t = re.sub(pat, rep, t)
    if k == 'animals':
        t = t.replace('Leo the Lion', 'Rory the Lion').replace('The Bear Who Loved Honey', 'Bobo the Bear Who Loved Honey')
    suffix = {'stories': 'Kids Stories', 'numbers': 'Learn Math for Kids', 'words': 'Kids Vocabulary', 'animals': 'Animal Stories for Kids',
              'science': 'Science for Kids', 'puzzles': 'Brain Games for Kids', 'morals': 'Moral Stories for Kids', 'songs': 'Kids Songs'}[k]
    return f'{t} {LAB[k][1]} | {suffix} | Tiny Sparks'

RIG = {'bear', 'panda', 'rabbit', 'bunny', 'cat', 'kitten', 'fox', 'lion', 'mouse', 'pig', 'owl', 'frog', 'monkey'}
NEW = ['elephant', 'penguin', 'dog', 'puppy', 'puppies', 'butterfly', 'butterflies', 'turtle', 'deer', 'wolf', 'giraffe', 'zebra', 'hippo',
       'crocodile', 'kangaroo', 'parrot', 'dolphin', 'whale', 'squirrel', 'firefly', 'bee', 'bees', 'ant', 'ants', 'peacock', 'crow', 'spider',
       'dinosaur', 'dinosaurs', 't-rex', 'bird', 'birds', 'fish', 'horse', 'goat', 'chicken', 'insects', 'robot', 'farm animals', 'ocean animals',
       'jungle', 'baby animal']
def rig_ready(k, idea):
    low = idea.lower()
    miss = sorted({w for w in NEW if re.search(r'(?<![a-z])' + re.escape(w) + r'(?![a-z])', low)})
    if miss: return 'Build first: ' + ', '.join(miss[:3])
    return 'Yes'

def age(k, i, idea):
    l = idea.lower()
    hard = ['multiplication', 'division', 'fraction', 'money', 'black hole', 'artificial intelligence', 'internet', 'computers', 'electricity', 'batteries',
            'galaxy', 'earthquake', 'hurricane', 'tornado', 'gravity', 'logic', 'detective', '20-question', 'championship', 'digital safety', 'history of wheels']
    mid = ['skip counting', 'odd', 'even', 'greater', 'addition', 'subtraction', 'pattern', '1–50', '1–100', 'clock', 'months', 'verbs', 'expressions',
           'questions', 'quiz', 'riddle', 'memory', 'maze', 'differences', 'volcano', 'solar system', 'planet', 'mercury', 'venus', 'earth', 'mars', 'jupiter',
           'saturn', 'uranus', 'neptune', 'space', 'rocket', 'astronaut', 'human body', 'brain', 'heart', 'stomach', 'bones', 'muscles', 'fossil', 'dinosaur',
           'cycle', 'magnet', 'experiment', 'science', 'engineer', 'bridge', 'map', 'inventor', 'king', 'prince', 'knight', 'kingdom', 'castle', 'cultures', 'languages', 'time']
    if any(w in l for w in hard): return '9–12' if any(w in l for w in ['black hole', 'artificial intelligence', 'internet', 'electricity']) else '6–8'
    if any(w in l for w in mid): return '6–8'
    return '3–5'

def fmt(k, idea):
    l = idea.lower()
    longw = ['ultimate', 'quiz', 'challenge', 'championship', '1–50', '1–100', 'a–z', 'numbers 1–20', 'adventure', 'mystery', 'olympics', 'contest',
             'school:', 'treasure hunt', 'five ', 'solar system', 'life cycle', 'water cycle', 'memory game', 'matching game']
    if k == 'songs': return 'Long 16:9 + Short cut', '2–3 min (Short: 45 s chorus)'
    if k in ('stories', 'morals', 'animals'): return 'Long 16:9 + Short teaser', '3–5 min (Short: 50 s)'
    if any(w in l for w in longw) or re.match(r'^\d+ ', idea): return 'Long 16:9', '4–8 min'
    return 'Short 9:16', '45–60 s'

def cast(k, title):
    base = {'stories': 'Mila, Leo, Bobo + narrator', 'numbers': 'Bobo + Mila (counting fingers)', 'words': 'Mila + French narrator',
            'animals': 'guest animal(s) + Bobo', 'science': 'Prof. Hoot + Leo (wrong guess) + Mila', 'puzzles': 'Prof. Hoot + Leo + score strip',
            'morals': 'narrator + guest animals / Mila & Leo', 'songs': 'Mila sings, all dance'}[k]
    names = [n for n in ('Mila', 'Leo', 'Bobo') if re.search(r'\b' + n + r'\b', title)]
    return (', '.join(names) + ' (lead) · ' + base) if names and k == 'stories' else base

VALUES = [('truth', 'honesty'), ('honest', 'honesty'), ('share', 'sharing'), ('shared', 'sharing'), ('greedy', 'sharing'), ('selfish', 'sharing'), ('kind', 'kindness'),
          ('humility', 'humility'), ('proud', 'humility'), ('best', 'humility'), ('attention', 'humility'), ('gave up', 'perseverance'), ('give up', 'perseverance'),
          ('tried', 'perseverance'), ('failure', 'perseverance'), ('afraid', 'courage'), ('brave', 'courage'), ('believed', 'self-belief'), ('sorry', 'saying sorry'),
          ('promise', 'keeping promises'), ('please', 'politeness'), ('thank', 'gratitude'), ('listen', 'listening'), ('patien', 'patience'), ('angry', 'managing anger'),
          ('lazy', 'hard work'), ('hardworking', 'hard work'), ('jealous', 'jealousy'), ('help', 'helping others'), ('water', 'caring for the planet'),
          ('tree', 'caring for the planet'), ('forest', 'caring for the planet'), ('ocean', 'caring for the planet'), ('recycle', 'caring for the planet'),
          ('river', 'caring for the planet'), ('park', 'caring for the planet'), ('seed', 'caring for the planet'), ('left behind', 'including others'),
          ('new student', 'including others'), ('lonely', 'friendship'), ('friend', 'friendship'), ('forgave', 'forgiveness'), ('feelings', 'empathy'),
          ('empathy', 'empathy'), ('family', 'family'), ('treasure', 'what really matters'), ('gold', 'what really matters'), ('wish', 'what really matters'),
          ('together', 'teamwork'), ('race', 'friendship over winning'), ('competition', 'friendship over winning'), ('gift', 'generosity'), ('clever', 'thinking smart')]
def objective(k, idea):
    l = idea.lower()
    if k in ('morals',):
        for kw, v in VALUES:
            if kw in l: return f'Value: {v}'
        return 'Value: kindness'
    if k == 'stories':
        m = re.search(r'(?:Learns?|Learned|Discovers?)\s+(?:to |About |Why |How |the |From )?(.+)', idea)
        if m: return 'Life skill: ' + m.group(1).rstrip('.')
        for kw, v in VALUES:
            if kw in l: return f'Life skill: {v}'
        return 'Curiosity + problem solving'
    if k == 'numbers': return 'Math: ' + idea
    if k == 'words': return 'Vocabulary EN→FR: ' + idea.replace(' English/French', '').replace(' in English and French', '')
    if k == 'animals': return 'Story + 3 true animal facts; value: ' + next((v for kw, v in VALUES if kw in l), 'friendship')
    if k == 'science': return 'Answer simply: ' + idea
    if k == 'puzzles':
        t = ('observation' if 'differ' in l or 'hidden' in l or 'find' in l else 'memory' if 'memory' in l or 'remember' in l
             else 'patterns' if 'pattern' in l or 'next' in l or 'before' in l else 'logic' if 'logic' in l or 'who ' in l or 'detective' in l
             else 'spatial thinking' if 'maze' in l or 'path' in l or 'help the' in l else 'matching/classifying' if 'match' in l or 'belong' in l or 'shadow' in l
             else 'reasoning')
        return 'Thinking skill: ' + t
    return 'Sing & learn: ' + idea.replace(' Song', '')

F = 'Arial'
def font(**k): return Font(name=F, **{'size': 10, **k})
BLUE = font(color='0000FF'); HEAD = font(bold=True, color='FFFFFF'); BOLD = font(bold=True)
H_FILL = PatternFill('solid', fgColor='3B2A4A'); Y = PatternFill('solid', fgColor='FFFF00')
thin = Side(style='thin', color='D9D9D9'); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
STATUS = ['Idea', 'Scripted', 'Rendered', 'Scheduled', 'Published']
SC = {'Idea': 'FFFFFF', 'Scripted': 'FFF3BF', 'Rendered': 'D0EBFF', 'Scheduled': 'E5DBFF', 'Published': 'D3F9D8'}
NCOL = {'stories': 'FFE8CC', 'numbers': 'FFE3E3', 'words': 'E7F5FF', 'animals': 'EBFBEE', 'science': 'EDF2FF', 'puzzles': 'FFF9DB', 'morals': 'F8F0FC', 'songs': 'FFF0F6'}

wb = Workbook()
s = wb.active; s.title = 'Start Here'
s['A1'] = 'Tiny Sparks — 800-video content calendar'; s['A1'].font = font(bold=True, size=16)
s['A2'] = 'One video a day for 800 days (~2.2 years). Edit only the YELLOW cells and the blue columns; black cells are calculated.'; s['A2'].font = font(italic=True)
s['A4'], s['B4'], s['C4'] = 'Setting', 'Value', 'What it does'
for c in 'ABC': s[f'{c}4'].font = HEAD; s[f'{c}4'].fill = H_FILL
s['A5'] = 'Launch date (Day 1, a Monday)'; s['B5'] = dt.date(2026, 10, 19); s['B5'].number_format = 'ddd dd mmm yyyy'
s['C5'] = 'Every date in the Calendar = launch date + (Day # − 1). Keep it a Monday so the theme days stay on their weekdays.'
s['A6'] = 'Hours from ET to YOUR time'; s['B6'] = 5
s['C6'] = 'Times are planned in US Eastern (ET), the biggest English kids audience. Morocco/Western Europe ≈ +5 to +6 (US clocks change 1 Nov and in March).'
for c in ('B5', 'B6'): s[c].fill = Y; s[c].font = BLUE
s['A8'] = 'Post time by weekday (ET)'; s['A8'].font = BOLD
s['A9'], s['B9'], s['C9'] = 'Weekday', 'Time (ET)', 'Why'
for c in 'ABC': s[f'{c}9'].font = HEAD; s[f'{c}9'].fill = H_FILL
times = [('Monday', 16, 'after school / after nap — kids + parent co-viewing'), ('Tuesday', 16, ''), ('Wednesday', 16, ''), ('Thursday', 16, ''),
         ('Friday', 16, 'start of the weekend binge'), ('Saturday', 9, 'weekend morning: the biggest kids-viewing window'), ('Sunday', 9, 'weekend morning (songs)')]
for i, (d_, h, why) in enumerate(times, start=10):
    s[f'A{i}'] = d_; s[f'B{i}'] = dt.time(h, 0); s[f'B{i}'].number_format = 'hh:mm'; s[f'B{i}'].fill = Y; s[f'B{i}'].font = BLUE; s[f'C{i}'] = why
s['A18'] = 'Weekly rotation'; s['A18'].font = BOLD
rot = [('Monday', 'Story Monday: 🧠 Educational Story (odd weeks) / 📚 Moral Story (even weeks)'), ('Tuesday', 'Math Tuesday: 🔢 Numbers & Math'),
       ('Wednesday', 'Word Wednesday: 🔤 English/French'), ('Thursday', 'Animal Thursday: 🦁 Animal Story'), ('Friday', 'Science Friday: 🚀 Science'),
       ('Saturday', 'Puzzle Saturday: 🧩 Puzzle'), ('Sunday', 'Song Sunday: 🎵 Song'),
       ('After week 100', 'Story Season (days 701–800): the remaining 50 educational + 50 moral stories, alternating daily')]
for i, (a, b) in enumerate(rot, start=19): s[f'A{i}'] = a; s[f'B{i}'] = b
s['A28'] = 'Status'; s['A28'].font = BOLD
for i, st in enumerate(STATUS, start=29):
    s[f'A{i}'] = st; s[f'A{i}'].fill = PatternFill('solid', fgColor=SC[st])
    s[f'B{i}'] = {'Idea': 'not started', 'Scripted': 'script.md + beats.json done', 'Rendered': 'final MP4 exported',
                  'Scheduled': 'uploaded + scheduled in YouTube Studio', 'Published': 'live: paste the link, add views after 48 h and 7 days'}[st]
s['A35'] = 'Columns in the Calendar'; s['A35'].font = BOLD
notes = [('YouTube title', 'the idea adapted to the locked cast (Tommy/Max/Alex → Leo, Mia/Emma/Lily/Sara → Mila, Benny → Bobo) + niche tag + channel'),
         ('Age', '3–5 / 6–8 / 9–12, estimated from the topic; adjust freely'), ('Format / Length', 'Short 9:16 vs long 16:9 (+ a Short cut), from the niche skill rules'),
         ('Rig ready?', '"Yes" = every animal in it already exists in the Critter rig; otherwise the species to add to critter.tsx first'),
         ('Priority', 'P1 = first 13 weeks, P2 = rest of year 1, P3 = later. A "Build first" in the first 13 weeks is the to-do list for the art kit.')]
for i, (a, b) in enumerate(notes, start=36): s[f'A{i}'] = a; s[f'B{i}'] = b
s.column_dimensions['A'].width = 34; s.column_dimensions['B'].width = 30; s.column_dimensions['C'].width = 100
for row in s.iter_rows():
    for c in row:
        if c.value is not None and c.font.name != F: c.font = font()
for c in ('A1',): s[c].font = font(bold=True, size=16)

LAUNCH = "'Start Here'!$B$5"; OFF = "'Start Here'!$B$6"; TT = "'Start Here'!$B$10:$B$16"

cal = wb.create_sheet('Calendar')
cols = [('Day #', 7), ('Week', 6), ('Date', 15), ('Day', 6), ('Theme day', 15), ('Niche', 22), ('Niche #', 7), ('Idea (original)', 44), ('YouTube title', 66),
        ('Age', 7), ('Format', 22), ('Length', 24), ('Characters', 34), ('Learning objective', 44), ('Rig ready?', 26), ('Priority', 8), ('Status', 11),
        ('Post time (ET)', 10), ('Your time', 9), ('YouTube link', 30), ('Views 48h', 10), ('Views 7 days', 11), ('Subs gained', 10), ('Avg % viewed', 10), ('Notes', 36)]
for i, (h, w) in enumerate(cols, start=1):
    c = cal.cell(1, i, h); c.font = HEAD; c.fill = H_FILL; c.alignment = Alignment(wrap_text=True, vertical='center')
    cal.column_dimensions[get_column_letter(i)].width = w
cal.row_dimensions[1].height = 30; cal.freeze_panes = 'I2'; cal.auto_filter.ref = f'A1:Y{len(slots) + 1}'
for r, (d_, theme, k, idx) in enumerate(slots, start=2):
    idea = IDEAS[k][idx]; title = yt_title(k, idea); f_, ln = fmt(k, idea); rr = rig_ready(k, idea)
    pr = 'P1' if d_ <= 91 else 'P2' if d_ <= 365 else 'P3'
    vals = {1: d_, 2: f'=INT((A{r}-1)/7)+1', 3: f'={LAUNCH}+A{r}-1', 4: f'=TEXT(C{r},"ddd")', 5: theme, 6: f'{LAB[k][1]} {LAB[k][0]}', 7: idx + 1,
            8: idea, 9: title, 10: age(k, idx, idea), 11: f_, 12: ln, 13: cast(k, title), 14: objective(k, idea), 15: rr, 16: pr, 17: 'Idea',
            18: f'=INDEX({TT},WEEKDAY(C{r},2))', 19: f'=MOD(R{r}+{OFF}/24,1)'}
    for c, v in vals.items(): cal.cell(r, c, v)
    cal.cell(r, 3).number_format = 'ddd dd mmm yyyy'; cal.cell(r, 18).number_format = 'hh:mm'; cal.cell(r, 19).number_format = 'hh:mm'; cal.cell(r, 24).number_format = '0%'
    for c in range(1, 26):
        cell = cal.cell(r, c); cell.border = BOX
        cell.font = font() if c in (2, 3, 4, 18, 19) else BLUE
    cal.cell(r, 6).fill = PatternFill('solid', fgColor=NCOL[k])
LAST = len(slots) + 1
cal.cell(2, 25, 'Example: Made for Kids = YES · playlist "Kids Stories"').font = BLUE
dv = DataValidation(type='list', formula1='"' + ','.join(STATUS) + '"', allow_blank=True); cal.add_data_validation(dv); dv.add(f'Q2:Q{LAST}')
for st in STATUS[1:]:
    cal.conditional_formatting.add(f'Q2:Q{LAST}', FormulaRule(formula=[f'$Q2="{st}"'], fill=PatternFill('solid', fgColor=SC[st])))
cal.conditional_formatting.add(f'A2:E{LAST}', FormulaRule(formula=['$C2=TODAY()'], fill=PatternFill('solid', fgColor='FFE066')))
cal['O1'].comment = Comment('Species the Critter rig does not have yet. Add them to remotion/src/lib/kids/critter.tsx (and the model sheet) before that video.', 'Claude')
cal['J1'].comment = Comment('Estimated from the topic words; the channel targets 3–7, so 9–12 ideas are candidates to simplify or skip.', 'Claude')

# ── Niche Progress ──
p = wb.create_sheet('Niche Progress')
hd = ['Niche', 'Skill', 'Videos', 'Scripted', 'Rendered', 'Scheduled', 'Published', '% published', 'First post', 'Last post', 'Need new species']
for i, h in enumerate(hd, start=1):
    c = p.cell(1, i, h); c.font = HEAD; c.fill = H_FILL; p.column_dimensions[get_column_letter(i)].width = [26, 14, 8, 9, 9, 10, 10, 11, 15, 15, 14][i - 1]
for r, (k, l, e, sk) in enumerate(NICHES, start=2):
    name = f'{e} {l}'
    p.cell(r, 1, name).fill = PatternFill('solid', fgColor=NCOL[k]); p.cell(r, 2, sk)
    rng = f"Calendar!$F$2:$F${LAST},$A{r}"
    p.cell(r, 3, f'=COUNTIFS({rng})')
    for j, st in enumerate(['Scripted', 'Rendered', 'Scheduled', 'Published'], start=4):
        p.cell(r, j, f'=COUNTIFS({rng},Calendar!$Q$2:$Q${LAST},"{st}")')
    p.cell(r, 8, f'=IF(C{r}=0,0,G{r}/C{r})').number_format = '0%'
    p.cell(r, 9, f'=_xlfn.MINIFS(Calendar!$C$2:$C${LAST},Calendar!$F$2:$F${LAST},$A{r})').number_format = 'dd mmm yyyy'
    p.cell(r, 10, f'=_xlfn.MAXIFS(Calendar!$C$2:$C${LAST},Calendar!$F$2:$F${LAST},$A{r})').number_format = 'dd mmm yyyy'
    p.cell(r, 11, f'=COUNTIFS({rng},Calendar!$O$2:$O${LAST},"Build first*")')
    for c in range(1, 12): p.cell(r, c).font = font(); p.cell(r, c).border = BOX
p.cell(10, 1, 'Total').font = BOLD
for c in range(3, 8): L = get_column_letter(c); p.cell(10, c, f'=SUM({L}2:{L}9)').font = BOLD
p.cell(10, 8, '=IF(C10=0,0,G10/C10)').number_format = '0%'; p.cell(10, 11, '=SUM(K2:K9)').font = BOLD

# ── Weekly Stats (first 52 weeks) ──
w = wb.create_sheet('Weekly Stats')
hd = ['Week', 'Week starts', 'Planned', 'Published', 'Views (7 days)', 'Subs gained', 'Avg % viewed']
for i, h in enumerate(hd, start=1):
    c = w.cell(1, i, h); c.font = HEAD; c.fill = H_FILL; w.column_dimensions[get_column_letter(i)].width = [7, 15, 9, 10, 14, 12, 13][i - 1]
for i in range(52):
    r = i + 2
    w.cell(r, 1, i + 1); w.cell(r, 2, f'={LAUNCH}+{i * 7}').number_format = 'ddd dd mmm yyyy'
    rng = f'Calendar!$B$2:$B${LAST},$A{r}'
    w.cell(r, 3, f'=COUNTIFS({rng})'); w.cell(r, 4, f'=COUNTIFS({rng},Calendar!$Q$2:$Q${LAST},"Published")')
    w.cell(r, 5, f'=SUMIFS(Calendar!$V$2:$V${LAST},{rng})').number_format = '#,##0;-#,##0;-'
    w.cell(r, 6, f'=SUMIFS(Calendar!$W$2:$W${LAST},{rng})').number_format = '#,##0;-#,##0;-'
    w.cell(r, 7, f'=IFERROR(AVERAGEIFS(Calendar!$X$2:$X${LAST},{rng},Calendar!$X$2:$X${LAST},">0"),0)').number_format = '0%;-0%;-'
    for c in range(1, 8): w.cell(r, c).font = font(); w.cell(r, c).border = BOX
w.cell(54, 1, 'Total').font = BOLD
for c, L in ((3, 'C'), (4, 'D'), (5, 'E'), (6, 'F')):
    w.cell(54, c, f'=SUM({L}2:{L}53)').font = BOLD; w.cell(54, c).number_format = '#,##0;-#,##0;-'
w['A56'] = 'Calculated from the Calendar: set Status = Published and type the views / subs / avg % viewed there.'; w['A56'].font = font(italic=True)
w.freeze_panes = 'A2'

# ── Posting Guide ──
g = wb.create_sheet('Posting Guide')
g.column_dimensions['A'].width = 26; g.column_dimensions['B'].width = 120
g['A1'] = 'Posting Guide — Tiny Sparks'; g['A1'].font = font(bold=True, size=14)
rows = [
    ('Cadence', 'One video a day on fixed theme days (see Start Here). Consistency beats volume: never skip a theme day. If daily is too much, publish Mon/Wed/Fri/Sun and keep the same order — just change Day # spacing.'),
    ('Batching', 'Produce a full week (7 videos, one per niche) in one session and schedule them all in YouTube Studio (Status = Scheduled). Keep a buffer of 2 weeks ahead.'),
    ('Shorts + long', 'Stories, morals, animal stories and songs are long 16:9 episodes + a 45–60 s Short cut posted the same day at 18:00 ET. Math, words, science and puzzles are mostly Shorts; their "Ultimate / Quiz / 20…" ideas are long.'),
    ('Compilations', 'Every 4 weeks, stitch the month\'s long episodes of one niche into a 20–30 min compilation and post it on a Saturday at 9:00 ET (extra slot) — the watch-time engine of kids channels.'),
    ('Every upload', 'Made for Kids = YES · title from the Calendar · description = what kids learn (2 lines) + ages + lyrics for songs · playlist per niche · thumbnail (long videos): one big face + 1–3 words · Altered/synthetic content = No.'),
    ('After upload', 'Paste the link · after 48 h and 7 days type the views · Avg % viewed (Shorts > 70 %, long > 40 % is good) · note what worked.'),
    ('Times', 'Weekdays 16:00 ET (after school / nap), weekends 09:00 ET (morning). These are starting points; after ~4 weeks use YouTube Analytics → Audience → "When your viewers are on YouTube" and update the time table on Start Here.'),
    ('Quality bar', 'Original characters, original scripts, original songs (tools/gen_song.py), real teaching value and genuine variation in every video. YouTube demonetises mass-produced, near-identical kids videos (inauthentic content) — the niche skills enforce the variety.'),
    ('Language', 'English is the primary language (US audience); French is the bonus line ("Learn 20 Animals in English & French"). A full French channel can come later by re-voicing the same videos.'),
    ('Art to-do', 'Filter the Calendar column "Rig ready?" ≠ Yes within Priority P1: that is the list of animal species to add to the Critter rig before those days.'),
]
for i, (a, b) in enumerate(rows, start=3):
    g[f'A{i}'] = a; g[f'B{i}'] = b; g[f'A{i}'].font = BOLD; g[f'B{i}'].font = font()
    g[f'B{i}'].alignment = Alignment(wrap_text=True, vertical='top'); g[f'A{i}'].alignment = Alignment(vertical='top'); g.row_dimensions[i].height = 44

for sh in (s, g): sh.sheet_view.showGridLines = False
wb.calculation.fullCalcOnLoad = True
out = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'TINY-SPARKS-CALENDAR.xlsx')
if os.path.exists(out) and '--force' not in sys.argv:
    sys.exit(f'{out} exists — re-generating would erase statuses/stats. Pass --force to overwrite.')
wb.save(out); print('saved', out, LAST)
