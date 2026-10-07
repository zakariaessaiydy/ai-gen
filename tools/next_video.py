#!/usr/bin/env python3
"""
next_video.py — the token-cheap tracker for the Tiny Sparks 800-video calendar.

Claude never reads TINY-SPARKS-CALENDAR.xlsx (800 rows = thousands of tokens). This script reads it
and prints ONLY what is needed (a ~10-line brief), and writes the status back.

  python tools/next_video.py              # brief for the NEXT video to make (first row not made yet)
  python tools/next_video.py --done 12    # mark Day 12 as made (Status=Rendered, date + file in Notes)
       [--file kids-shorts/.../output/x.mp4]
  python tools/next_video.py --published 12 --link https://youtu.be/...   # after uploading
  python tools/next_video.py --status     # one-line progress summary
  python tools/next_video.py --peek 5     # the next 5 topics, one line each

"Made" = Status in Rendered / Scheduled / Published. Needs openpyxl.
"""
import argparse
import datetime as dt
import os
import re
import sys

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.environ.get("CALENDAR_XLSX", os.path.join(ROOT, "TINY-SPARKS-CALENDAR.xlsx"))
MADE = {"Rendered", "Scheduled", "Published"}
# Calendar columns (1-based)
C = dict(day=1, theme=5, niche=6, nidx=7, idea=8, title=9, age=10, fmt=11, length=12, cast=13, goal=14, rig=15, prio=16, status=17, link=20, notes=25)
SKILL = {"Educational Stories": "kids-stories", "Numbers & Math": "kids-numbers", "English/French Words": "kids-words", "Animal Stories": "kids-animals",
         "Science for Kids": "kids-science", "Puzzles & Problem-Solving": "kids-puzzles", "Moral Stories": "kids-morals", "Educational Songs": "kids-songs"}


def load():
    wb = load_workbook(XLSX)
    return wb, wb["Calendar"]


def launch(wb):
    v = wb["Start Here"]["B5"].value
    return v.date() if isinstance(v, dt.datetime) else v


def row_of(ws, day):
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, C["day"]).value == day:
            return r
    sys.exit(f"Day {day} not found")


def niche_name(cell):
    return re.sub(r"^[^A-Za-z]+", "", cell or "").strip()


def slug(s):
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")[:40]


def brief(wb, ws, r):
    g = lambda k: ws.cell(r, C[k]).value
    day = g("day")
    date = launch(wb) + dt.timedelta(days=day - 1)
    niche = niche_name(g("niche"))
    sk = SKILL.get(niche, "make-kids")
    nkey = sk.replace("kids-", "")
    proj = f"kids-shorts/tiny-sparks/{nkey}/day-{day:03d}-{slug(str(g('title')).split(' | ')[0])}"
    comp = f"Kids{nkey.capitalize()}{day:03d}"
    print(f"NEXT VIDEO — Day {day} · {date:%a %d %b %Y} · {g('theme')} · {g('prio')}")
    print(f"  niche:   {g('niche')} #{g('nidx')}  → skill: {sk} (read make-kids first)")
    print(f"  idea:    {g('idea')}")
    print(f"  title:   {g('title')}")
    print(f"  format:  {g('fmt')} · {g('length')} · ages {g('age')}")
    print(f"  cast:    {g('cast')}")
    print(f"  goal:    {g('goal')}")
    print(f"  rig:     {g('rig')}")
    print(f"  project: {proj}/   composition: {comp} (remotion/src/shots/kids-{nkey}-{day:03d}/)")
    print(f"  when done: python tools/next_video.py --done {day} --file {proj}/output/<file>.mp4")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--done", type=int)
    ap.add_argument("--file")
    ap.add_argument("--published", type=int)
    ap.add_argument("--link")
    ap.add_argument("--status", action="store_true")
    ap.add_argument("--peek", type=int)
    a = ap.parse_args()
    wb, ws = load()
    rows = range(2, ws.max_row + 1)

    if a.done or a.published:
        day = a.done or a.published
        r = row_of(ws, day)
        today = dt.date.today().isoformat()
        if a.done:
            ws.cell(r, C["status"]).value = "Rendered"
            note = f"made {today}" + (f" → {a.file}" if a.file else "")
        else:
            ws.cell(r, C["status"]).value = "Published"
            note = f"published {today}"
            if a.link:
                ws.cell(r, C["link"]).value = a.link
        old = ws.cell(r, C["notes"]).value
        ws.cell(r, C["notes"]).value = f"{old} · {note}" if old and not str(old).startswith("Example") else note
        wb.save(XLSX)
        made = sum(1 for x in rows if ws.cell(x, C["status"]).value in MADE)
        print(f"Day {day} → {ws.cell(r, C['status']).value}: {ws.cell(r, C['idea']).value}  ({made}/800 made)")
        return

    if a.status:
        made = sum(1 for x in rows if ws.cell(x, C["status"]).value in MADE)
        pub = sum(1 for x in rows if ws.cell(x, C["status"]).value == "Published")
        nxt = next((ws.cell(x, C["day"]).value for x in rows if ws.cell(x, C["status"]).value not in MADE), None)
        today_day = (dt.date.today() - launch(wb)).days + 1
        if today_day < 1:
            when = f"launch in {1 - today_day} day(s)"
        elif not nxt:
            when = "all done"
        else:
            when = f"today = Day {today_day} · " + (f"{today_day - nxt} day(s) behind" if today_day > nxt else f"{nxt - today_day} day(s) ahead")
        print(f"{made}/800 made · {pub} published · next = Day {nxt} · {when}")
        return

    if a.peek:
        n = 0
        for x in rows:
            if ws.cell(x, C["status"]).value not in MADE:
                print(f"Day {ws.cell(x, C['day']).value:3d} · {ws.cell(x, C['niche']).value} · {ws.cell(x, C['idea']).value}")
                n += 1
                if n >= a.peek:
                    break
        return

    r = next((x for x in rows if ws.cell(x, C["status"]).value not in MADE), None)
    if r is None:
        print("All 800 videos are made 🎉")
        return
    brief(wb, ws, r)


if __name__ == "__main__":
    main()
